"""Automated test suite for Document Processing & Text Extraction (Phase 5)."""

import io
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from typing import Any, Dict, List, Optional
from bson import ObjectId
import docx
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pptx import Presentation
import pymupdf

from app.config import get_settings
from app.core.security import create_access_token, hash_password
from app.ingestion.ocr_processor import OCRError
from app.main import app
from app.models.course import CourseInDB
from app.models.material import MaterialInDB, ProcessingStatus, SourceType
from app.repositories.extracted_content_repository import ExtractedContentRepository
from app.repositories.material_repository import MaterialRepository


class InMemoryExtractedContentCollection:
    """Mock in-memory collection simulating MongoDB extracted_content collection."""

    def __init__(self) -> None:
        self.docs: Dict[str, Dict[str, Any]] = {}

    def insert_many(self, docs: List[Dict[str, Any]]) -> Any:
        inserted_ids = []
        for doc in docs:
            doc_copy = dict(doc)
            _id = ObjectId()
            doc_copy["_id"] = _id
            self.docs[str(_id)] = doc_copy
            inserted_ids.append(_id)

        class InsertManyResult:
            def __init__(self, ids: List[ObjectId]):
                self.inserted_ids = ids

        return InsertManyResult(inserted_ids)

    def find(self, query: Dict[str, Any] = None) -> Any:
        items = list(self.docs.values())
        if query and "material_id" in query:
            mat_id = query["material_id"]
            items = [d for d in items if d.get("material_id") == mat_id]

        class Cursor:
            def __init__(self, items: List[Dict[str, Any]]):
                self.items = [dict(i) for i in items]

            def sort(self, key: str, direction: int = 1) -> Any:
                self.items.sort(key=lambda x: x.get(key, 0), reverse=(direction == -1))
                return self

            def __iter__(self) -> Any:
                return iter(self.items)

        return Cursor(items)

    def delete_many(self, query: Dict[str, Any]) -> Any:
        deleted = 0
        if "material_id" in query:
            mat_id = query["material_id"]
            for k, d in list(self.docs.items()):
                if d.get("material_id") == mat_id:
                    del self.docs[k]
                    deleted += 1

        class DeleteResult:
            def __init__(self, count: int):
                self.deleted_count = count

        return DeleteResult(deleted)


class InMemoryMaterialsCollection:
    """Mock in-memory collection simulating MongoDB materials collection."""

    def __init__(self) -> None:
        self.docs: Dict[str, Dict[str, Any]] = {}

    def insert_one(self, doc: Dict[str, Any]) -> Any:
        doc_copy = dict(doc)
        _id = ObjectId()
        doc_copy["_id"] = _id
        id_str = str(_id)
        self.docs[id_str] = doc_copy

        class InsertResult:
            def __init__(self, inserted_id: ObjectId):
                self.inserted_id = inserted_id

        return InsertResult(_id)

    def find_one(self, query: Dict[str, Any]) -> Optional[Dict[str, Any]]:
        if "_id" in query:
            _id = query["_id"]
            for d in self.docs.values():
                if d.get("_id") == _id:
                    return dict(d)
        if "course_code" in query and "original_filename" in query:
            code = query["course_code"]
            fn = query["original_filename"]
            for d in self.docs.values():
                if d.get("course_code") == code and d.get("original_filename") == fn:
                    return dict(d)
        return None

    def find(self, query: Dict[str, Any] = None) -> Any:
        items = list(self.docs.values())
        if query and "course_code" in query:
            code = query["course_code"]
            items = [d for d in items if d.get("course_code") == code]

        class Cursor:
            def __init__(self, items: List[Dict[str, Any]]):
                self.items = [dict(i) for i in items]

            def skip(self, n: int) -> Any:
                self.items = self.items[n:]
                return self

            def limit(self, n: int) -> Any:
                self.items = self.items[:n]
                return self

            def sort(self, key: str, direction: int = 1) -> Any:
                self.items.sort(key=lambda x: x.get(key, ""), reverse=(direction == -1))
                return self

            def __iter__(self) -> Any:
                return iter(self.items)

        return Cursor(items)

    def update_one(self, query: Dict[str, Any], update: Dict[str, Any]) -> Any:
        modified = 0
        if "_id" in query:
            _id = query["_id"]
            for d in self.docs.values():
                if d.get("_id") == _id:
                    if "$set" in update:
                        d.update(update["$set"])
                    modified = 1

        class UpdateResult:
            def __init__(self, count: int):
                self.modified_count = count

        return UpdateResult(modified)

    def delete_one(self, query: Dict[str, Any]) -> Any:
        deleted = 0
        if "_id" in query:
            _id = query["_id"]
            for k, d in list(self.docs.items()):
                if d.get("_id") == _id:
                    del self.docs[k]
                    deleted = 1

        class DeleteResult:
            def __init__(self, count: int):
                self.deleted_count = count

        return DeleteResult(deleted)


class TestDocumentProcessing(unittest.TestCase):
    """Test suite covering Phase 5 Document Processing and Text Extraction."""

    @classmethod
    def setUpClass(cls) -> None:
        """Set up test admin credentials, developer token, and temporary test storage."""
        cls.test_username = "processing_admin"
        cls.test_password = "SecureProcessingPassword123!"
        cls.test_password_hash = hash_password(cls.test_password)

        settings = get_settings()
        settings.ADMIN_USERNAME = cls.test_username
        settings.ADMIN_PASSWORD_HASH = cls.test_password_hash

        cls.temp_storage = tempfile.mkdtemp(prefix="hots_proc_test_")
        settings.EXTERNAL_STORAGE_PATH = cls.temp_storage

        cls.dev_token = create_access_token(subject=cls.test_username, role="developer")
        cls.auth_headers = {"Authorization": f"Bearer {cls.dev_token}"}

    @classmethod
    def tearDownClass(cls) -> None:
        """Clean up temporary test files."""
        if os.path.exists(cls.temp_storage):
            shutil.rmtree(cls.temp_storage, ignore_errors=True)

    def setUp(self) -> None:
        """Patch database collections."""
        from unittest.mock import patch

        self.mock_materials_coll = InMemoryMaterialsCollection()
        self.mock_extracted_coll = InMemoryExtractedContentCollection()

        self.patcher_mat = patch(
            "app.repositories.material_repository.get_materials_collection",
            return_value=self.mock_materials_coll,
        )
        self.patcher_ext = patch(
            "app.repositories.extracted_content_repository.get_extracted_content_collection",
            return_value=self.mock_extracted_coll,
        )
        self.patcher_mat.start()
        self.patcher_ext.start()

        # Create Course in DB
        self.sample_course = CourseInDB(
            id=str(ObjectId()),
            course_code="BCSE301",
            course_name="Database Management Systems",
        )
        self.patcher_course = patch(
            "app.repositories.course_repository.CourseRepository.get_course",
            side_effect=lambda code: self.sample_course if code.strip().upper() == "BCSE301" else None,
        )
        self.patcher_course.start()

    def tearDown(self) -> None:
        """Stop patches."""
        self.patcher_mat.stop()
        self.patcher_ext.stop()
        self.patcher_course.stop()

    def _create_sample_pdf(self, path: Path, text: str) -> None:
        """Create a real digital test PDF with PyMuPDF."""
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 50), text, fontsize=12)
        doc.save(str(path))
        doc.close()

    def _create_sample_pptx(self, path: Path, title: str, content: str) -> None:
        """Create a real test PPTX with python-pptx."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = content
        prs.save(str(path))

    def _create_sample_docx(self, path: Path, text: str) -> None:
        """Create a real test DOCX with python-docx."""
        doc = docx.Document()
        doc.add_heading("Course Syllabus Heading", level=1)
        doc.add_paragraph(text)
        doc.save(str(path))

    def _create_sample_image(self, path: Path, text: str) -> None:
        """Create a real test image with PIL."""
        img = Image.new("RGB", (400, 100), color=(255, 255, 255))
        d = ImageDraw.Draw(img)
        d.text((10, 10), text, fill=(0, 0, 0))
        img.save(str(path))

    def test_1_process_unauthenticated_rejected(self) -> None:
        """Verify POST /dev/materials/{material_id}/process requires developer authentication."""
        with TestClient(app) as client:
            res = client.post("/dev/materials/6a9cf9648fe403384c67ab2a/process")
            self.assertEqual(res.status_code, 401)

    def test_2_process_nonexistent_material_returns_404(self) -> None:
        """Verify 404 is returned when material does not exist in MongoDB."""
        with TestClient(app) as client:
            res = client.post("/dev/materials/6a9cf9648fe403384c67ab2a/process", headers=self.auth_headers)
            self.assertEqual(res.status_code, 404)
            self.assertIn("not found", res.json()["detail"].lower())

    def test_3_process_missing_stored_file_returns_400_and_fails_status(self) -> None:
        """Verify missing disk file returns 400 and updates status to failed in MongoDB."""
        with TestClient(app) as client:
            # Upload initial record
            files = {"file": ("test.pdf", io.BytesIO(b"%PDF-1.4 header"), "application/pdf")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "syllabus"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]
            stored_name = up_res.json()["stored_filename"]

            # Delete physical file behind the scenes
            disk_file = Path(self.temp_storage) / "BCSE301" / stored_name
            if disk_file.exists():
                disk_file.unlink()

            # Process
            proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(proc_res.status_code, 400)
            self.assertIn("not found on disk", proc_res.json()["detail"].lower())

            # Verify DB state is failed
            mat = MaterialRepository.get_material(mat_id)
            self.assertIsNotNone(mat)
            self.assertEqual(mat.processing_status, ProcessingStatus.FAILED)
            self.assertIn("not found on disk", mat.error_message.lower())

    def test_4_process_digital_pdf_success(self) -> None:
        """Verify processing a valid digital PDF extracts text via PyMuPDF."""
        with TestClient(app) as client:
            # Create real PDF on disk
            pdf_bytes = io.BytesIO()
            doc = pymupdf.open()
            page = doc.new_page()
            page.insert_text((50, 50), "Introduction to Relational Databases and SQL Queries", fontsize=12)
            doc.save(pdf_bytes)
            doc.close()
            pdf_bytes.seek(0)

            files = {"file": ("dbms_syllabus.pdf", pdf_bytes, "application/pdf")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "syllabus"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            # Process
            proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(proc_res.status_code, 200)
            data = proc_res.json()
            self.assertEqual(data["material_id"], mat_id)
            self.assertEqual(data["processing_status"], "processed")
            self.assertEqual(data["pages_processed"], 1)
            self.assertEqual(data["extraction_method"], "pymupdf")
            self.assertTrue(data["total_characters"] > 20)

            # Check DB records
            mat = MaterialRepository.get_material(mat_id)
            self.assertEqual(mat.processing_status, ProcessingStatus.PROCESSED)

            extracted_pages = ExtractedContentRepository.get_content_by_material(mat_id)
            self.assertEqual(len(extracted_pages), 1)
            self.assertIn("Relational Databases", extracted_pages[0].text)

    def test_5_process_pptx_success(self) -> None:
        """Verify processing a PowerPoint presentation extracts slide text via python-pptx."""
        with TestClient(app) as client:
            pptx_bytes = io.BytesIO()
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "Unit 1: ER Modeling"
            slide.placeholders[1].text = "Entity Relationship Diagrams and Cardinality Constraints"
            prs.save(pptx_bytes)
            pptx_bytes.seek(0)

            files = {"file": ("lecture1.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "lecture_material"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(proc_res.status_code, 200)
            data = proc_res.json()
            self.assertEqual(data["processing_status"], "processed")
            self.assertEqual(data["extraction_method"], "python-pptx")
            self.assertEqual(data["pages_processed"], 1)

            extracted = ExtractedContentRepository.get_content_by_material(mat_id)
            self.assertEqual(len(extracted), 1)
            self.assertIn("ER Modeling", extracted[0].text)

    def test_6_process_docx_success(self) -> None:
        """Verify processing a Word document extracts text via python-docx."""
        with TestClient(app) as client:
            docx_bytes = io.BytesIO()
            doc = docx.Document()
            doc.add_heading("Normalization Chapter Notes", level=1)
            doc.add_paragraph("First Normal Form (1NF), 2NF, 3NF, and BCNF definitions.")
            doc.save(docx_bytes)
            docx_bytes.seek(0)

            files = {"file": ("notes.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "lecture_material"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(proc_res.status_code, 200)
            data = proc_res.json()
            self.assertEqual(data["processing_status"], "processed")
            self.assertEqual(data["extraction_method"], "python-docx")

            extracted = ExtractedContentRepository.get_content_by_material(mat_id)
            self.assertEqual(len(extracted), 1)
            self.assertIn("Normalization Chapter Notes", extracted[0].text)

    def test_7_process_image_ocr_with_mock(self) -> None:
        """Verify image processing triggers OCR pipeline and saves extracted text."""
        from unittest.mock import patch
        with TestClient(app) as client:
            img_bytes = io.BytesIO()
            img = Image.new("RGB", (300, 100), color=(255, 255, 255))
            img.save(img_bytes, format="PNG")
            img_bytes.seek(0)

            files = {"file": ("scan_note.png", img_bytes, "image/png")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "reference_book"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            # Mock OCR extraction to return deterministic recognized text
            with patch("app.ingestion.ocr_processor.OCRProcessor.extract_text_from_image", return_value="B-Tree Indexing and Hashing Techniques"):
                proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
                self.assertEqual(proc_res.status_code, 200)
                data = proc_res.json()
                self.assertEqual(data["extraction_method"], "tesseract_ocr")
                self.assertEqual(data["processing_status"], "processed")

                extracted = ExtractedContentRepository.get_content_by_material(mat_id)
                self.assertEqual(len(extracted), 1)
                self.assertEqual(extracted[0].text, "B-Tree Indexing and Hashing Techniques")

    def test_8_ocr_unavailable_failure_handling(self) -> None:
        """Verify that OCR failure (e.g. missing Tesseract) is handled gracefully without crashing."""
        from unittest.mock import patch
        with TestClient(app) as client:
            img_bytes = io.BytesIO()
            img = Image.new("RGB", (200, 50), color=(255, 255, 255))
            img.save(img_bytes, format="PNG")
            img_bytes.seek(0)

            files = {"file": ("scan_fail.png", img_bytes, "image/png")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "exam_paper"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            with patch("app.ingestion.ocr_processor.OCRProcessor.extract_text_from_image", side_effect=OCRError("Tesseract OCR binary not found")):
                proc_res = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
                self.assertEqual(proc_res.status_code, 500)
                self.assertIn("tesseract ocr binary not found", proc_res.json()["detail"].lower())

                # Status in DB should be failed
                mat = MaterialRepository.get_material(mat_id)
                self.assertEqual(mat.processing_status, ProcessingStatus.FAILED)

    def test_9_idempotent_reprocessing_cleans_duplicates(self) -> None:
        """Verify calling process multiple times on the same material cleans up old records and prevents duplicates."""
        with TestClient(app) as client:
            pdf_bytes = io.BytesIO()
            doc = pymupdf.open()
            page = doc.new_page()
            page.insert_text((50, 50), "ACID Properties: Atomicity, Consistency, Isolation, Durability", fontsize=12)
            doc.save(pdf_bytes)
            doc.close()
            pdf_bytes.seek(0)

            files = {"file": ("acid_notes.pdf", pdf_bytes, "application/pdf")}
            up_res = client.post("/dev/materials", data={"course_code": "BCSE301", "source_type": "lecture_material"}, files=files, headers=self.auth_headers)
            mat_id = up_res.json()["id"]

            # First process
            res1 = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(res1.status_code, 200)
            self.assertEqual(len(ExtractedContentRepository.get_content_by_material(mat_id)), 1)

            # Reprocess
            res2 = client.post(f"/dev/materials/{mat_id}/process", headers=self.auth_headers)
            self.assertEqual(res2.status_code, 200)
            # Must remain exactly 1 record, not 2
            self.assertEqual(len(ExtractedContentRepository.get_content_by_material(mat_id)), 1)


if __name__ == "__main__":
    unittest.main()
